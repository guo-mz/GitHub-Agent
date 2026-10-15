# -*- coding: utf-8 -*-
"""生成 GitHub Agent 面试材料 Word 文档。"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent


def set_run_font(run, *, size=11, bold=False, color=None, name="微软雅黑"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    if color:
        run.font.color.rgb = color
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size={1: 18, 2: 14, 3: 12}.get(level, 11), bold=True, name="微软雅黑")
    return p


def add_para(doc, text, *, bold=False, size=11, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    return p


def add_bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(it)
        set_run_font(run, size=11)
        p.paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                set_run_font(run, size=10, bold=True)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    set_run_font(run, size=9)
    doc.add_paragraph()


def style_doc(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)


def cover(doc, title, subtitle):
    style_doc(doc)
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(title)
    set_run_font(r, size=22, bold=True, color=RGBColor(0x1A, 0x36, 0x5D))
    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = s.add_run(subtitle)
    set_run_font(r2, size=12, color=RGBColor(0x4A, 0x55, 0x68))
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = note.add_run("依据 github-agent 工作区真实代码整理 · 勿夸大未验证能力")
    set_run_font(r3, size=10, color=RGBColor(0x71, 0x80, 0x96))
    doc.add_paragraph()


# ─── 文档 1：完整总册 ───────────────────────────────────────────


def build_full():
    doc = Document()
    cover(doc, "GitHub Agent 技术面试材料（完整版）", "15 分钟项目介绍 · PPT · 追问 · 演示 · 多时长版本")

    add_heading(doc, "〇、代码事实核查", 1)
    add_para(doc, "讲解前必须区分三种状态：已实现且有测试证据 / 代码存在但运行未充分验证 / 尚未实现或硬关闭。")

    add_heading(doc, "核查过的关键文件", 2)
    add_table(
        doc,
        ["领域", "关键路径"],
        [
            ["入口/路由", "backend/app/main.py、api/chat.py、api/approval.py"],
            ["状态图", "agent/graph.py、state.py、edges.py、nodes.py"],
            ["多 Agent", "agents/analysis|planning|coding|testing|publish_agent.py"],
            ["Checkpointer", "core/checkpointer.py、core/runtime_mode.py"],
            ["工作区/快照", "services/workspace_service.py"],
            ["写门禁", "write_gate.py、approval_gate.py、sandbox_service.py"],
            ["前端 SSE", "frontend/src/api/index.ts、stores/chat.ts"],
            ["测试", "tests/test_edges|checkpoint_hitl|publish_agent|security_phase*.py 等"],
            ["启动/文档", "scripts/start-all.ps1、demo.ps1、README.md、.env.example"],
        ],
    )

    add_heading(doc, "能力三态分类", 2)
    add_heading(doc, "已实现且有代码/测试证据", 3)
    add_bullets(
        doc,
        [
            "Vue 3 + FastAPI 单体；POST /api/chat/stream SSE；POST /api/approval/{tid}",
            "LangGraph StateGraph：多 Agent 主路径 + 旧 ReAct 旁路",
            "三道 HITL：await_plan_confirm / await_diff_approve / await_publish_approve（interrupt + Command(resume)）",
            "工作区 base/work/approved_snapshot、content_hash 冻结与校验",
            "write_gate.permit_github_write；evaluate_test_gate 始终拒绝（含伪造 passed=True）",
            "Checkpointer：AsyncRedisSaver 优先，失败降级 MemorySaver；无耐久 checkpoint 默认 503 拒新 HITL",
            "操作员 Token、审批 resume 校验、防重复审批",
        ],
    )
    add_heading(doc, "代码已实现，真实运行/集成尚未充分验证", 3)
    add_bullets(
        doc,
        [
            "DeepSeek 经 ChatOpenAI + LLM_BASE_URL（OpenAI 兼容）",
            "Redis Stack + AsyncRedisSaver 跨进程/跨重启 HITL",
            "前端完整浏览器联调、Docker Compose 全栈、微服务模式",
            "真实 GitHub 读操作（工具与 client 存在）",
        ],
    )
    add_heading(doc, "尚未实现或环境硬限制", 3)
    add_bullets(
        doc,
        [
            "可信隔离测试执行器：sandbox_status().can_execute 恒为 False",
            "经 Agent 管线真实发布 PR：被 evaluate_test_gate 硬关闭",
            "旧 ReAct 路径 tools_approved_node 对远程写已拒绝",
        ],
    )
    add_heading(doc, "README 与代码差异（面试别踩坑）", 2)
    add_bullets(
        doc,
        [
            "README 架构仍偏 ReAct；主写路径已是多 Agent 三道 HITL",
            "README 验收写「同意→真实 PR」；代码与安全节已硬关闭发布——以代码为准",
            "端口：start-all.ps1 常用 8001，README/demo.ps1 默认 8000",
        ],
    )

    # 口述稿
    add_heading(doc, "一、完整口述稿（约 15 分钟）", 1)

    add_heading(doc, "第一部分：项目背景与核心价值（约 1 分钟）", 2)
    add_para(
        doc,
        "我做的这个项目叫 GitHub Agent，本质是一个面向仓库改动的 AI Coding Agent Web 应用。"
        "场景很具体：日常改代码时，大模型可以帮你写补丁、提 PR，但如果只是「聊天框里吐一段代码」，"
        "中间缺三样东西——任务怎么拆、状态怎么管、高风险写操作谁拍板。"
        "一旦模型 hallucinate 了路径、漏了测试，或在 main 上直接推，代价很大。",
    )
    add_para(
        doc,
        "所以我把目标定成：不是做一个能聊天的代码助手，而是做一个能把"
        "「需求分析 → 规划 → 人工确认 → 本地工作区改代码 → Diff 审查 → 测试门禁 → 发布决策」"
        "串成状态机的系统。前端 Vue 3，后端 FastAPI，编排用 LangGraph，模型走 DeepSeek 的 OpenAI 兼容接口。"
        "需要强调一句：当前版本把「真实远程发布」故意硬关掉了——因为还没有可信的隔离测试执行器。"
        "我更愿意讲清楚「为什么不发」，而不是假装已经能安全发 PR。",
    )

    add_heading(doc, "第二部分：整体架构（约 2 分钟）", 2)
    add_para(doc, "整体可以看成四层：")
    add_bullets(
        doc,
        [
            "前端：任务输入、节点时间线、计划确认、Diff 查看、审批弹窗；用 fetch + ReadableStream 吃 SSE",
            "FastAPI：/api/chat/stream 开流；/api/approval/{tid} 提交决策并 resume；敏感接口挂操作员 Token",
            "LangGraph：AgentState 承载 phase、计划、Diff、哈希、测试报告、三道审批标志；compile 注入 Checkpointer",
            "外部：LLM（DeepSeek）、GitHub API、本地工作区、Redis（业务缓存 + 可选耐久 Checkpointer）",
        ],
    )
    add_para(
        doc,
        "调用链：用户发需求 → POST /api/chat/stream → assert_hitl_task_allowed → stream_agent → "
        "graph.astream_events → 事件映射 → 前端更新；interrupt 时审批 → Command(resume=decision) 继续。",
    )
    add_para(
        doc,
        "架构关系（口述）：Vue → FastAPI → LangGraph StateGraph → Checkpointer / LLM / Workspace / GitHub；"
        "另有 RedisService 管审批、幂等、限流（与 Checkpointer 不要混为一谈）。",
        bold=False,
    )

    add_heading(doc, "第三部分：核心业务流程（约 3 分钟）", 2)
    add_para(doc, "主路径（写类任务）按代码真实顺序：")
    add_bullets(
        doc,
        [
            "classify_intent：判定任务类型（read/write/review/qa）",
            "jev_decide：是否走规划链路",
            "repo_analysis → planning_agent：仓库分析与计划文档（LLM）",
            "await_plan_confirm：interrupt；确认前禁止编码",
            "coding_agent：只写隔离工作区 work/，产出真实 Diff，不写远程 GitHub",
            "await_diff_approve：展示前 freeze_approved_snapshot，批准绑 content_hash",
            "testing_agent：当前一律 not_run / unavailable",
            "await_publish_approve：再一道人工决策",
            "publish_agent：三门审批 + 测试门禁 + 哈希 + 快照 + 幂等 + WritePermit——但测试门禁今天恒失败，远程写走不到",
        ],
    )
    add_para(
        doc,
        "分工：LLM 负责理解与改码；条件边与门禁是确定性代码；计划/Diff/发布三处必须人拍板；"
        "任何「模型说 passed」都不能放行发布。纯 QA 可走旧 ReAct 旁路。",
    )

    add_heading(doc, "第四部分：关键技术设计（约 4 分钟）", 2)
    add_heading(doc, "1. 为什么选 LangGraph", 3)
    add_para(
        doc,
        "问题：流程长、有分支、有人工等待，普通脚本或 AgentExecutor 很难干净停住再继续。"
        "方案：StateGraph + 条件边 + Checkpointer。实现：build_graph()、route_after_*、compile(checkpointer=...)。"
        "收益：审批点是图的一等公民。限制：旧 checkpoint 兼容、节点命名（不能叫 plan）需小心。",
    )
    add_heading(doc, "2. Human-in-the-Loop", 3)
    add_para(
        doc,
        "节点内 interrupt(payload)；同 thread_id 审批；Command(resume=decision)。"
        "API 层 validate_approval_resume 防盲 resume、防已 resolved 复用。"
        "跨重启恢复依赖 Redis Stack 的 AsyncRedisSaver；Memory 模式仅进程内；"
        "无耐久 checkpoint 时 assert_hitl_task_allowed 挡虚假可恢复性。",
    )
    add_heading(doc, "3. FastAPI + SSE", 3)
    add_para(
        doc,
        "StreamingResponse + astream_events v2。断开时 request_cancel；冷启动瞬时故障可整轮 retry。"
        "断线后不是自动重连同一 SSE；靠任务状态 + 审批 resume。",
    )
    add_heading(doc, "4. Redis 双重角色", 3)
    add_para(
        doc,
        "LangGraph Checkpointer（AsyncRedisSaver 或 MemorySaver）≠ RedisService（审批/幂等/限流/会话）。"
        "把能不能跨重启恢复 interrupt 做成显式状态，写进 /health persistence。",
    )
    add_heading(doc, "5. 安全控制：为什么不信模型", 3)
    add_para(
        doc,
        "工作区隔离 → Diff 冻结快照 → SHA-256 content_hash → 发布只上传内存验证过的 approved_snapshot → "
        "evaluate_test_gate 对一切 passed 返回 untrusted_passed_no_secure_executor → "
        "只有 permit_github_write 上下文才能调写 API。"
        "诚实代价：端到端「同意就出 PR」目前做不到——这是设计选择。",
    )

    add_heading(doc, "第五部分：设计难点与解决思路（约 2 分钟）", 2)
    add_para(
        doc,
        "难点：LLM 不确定 + 远程写不可逆时，如何把「人批准的内容」钉死。"
        "若只存 Diff 文本或只信 work/，会出现人批 A、推上 B。"
        "因此 Diff 审批前 freeze_approved_snapshot；resume 校验 hash；Publish 再 verify_and_load_approved_snapshot；"
        "无隔离执行器则一律 not_run，门禁永不因模型声称通过而放行。"
        "收益：人工审批变成带完整性约束的安全协议；代价是发布链路暂时闭合。"
        "测试用 test_publish_agent、test_security_phase3 把拒绝路径钉死。",
    )

    add_heading(doc, "第六部分：项目现状与未来规划（约 1 分钟）", 2)
    add_bullets(
        doc,
        [
            "已完成：多 Agent 状态机、三道 HITL、工作区 Diff/哈希、写门禁、审批防重、SSE、发布拒绝路径",
            "未充分证明：真实 Redis Stack 跨重启、生产 DeepSeek/GitHub E2E、浏览器全链路",
            "明确未做：可信容器测试执行器 → 真实 Agent 发布 PR 硬关闭",
            "下一步：① 隔离执行器让门禁认执行器证明；② Redis Stack HITL 恢复做稳定验收；③ 任务观测与审计检索",
        ],
    )

    add_heading(doc, "第七部分：总结（约 1 分钟）", 2)
    add_para(
        doc,
        "这不是「调用大模型生成代码的聊天框」，而是一次把任务编排（LangGraph）、状态持久化（Checkpointer）、"
        "人工决策（interrupt/resume）、工程安全约束（快照哈希 + 写门禁 + 测试门禁）组合在一起的 Coding Agent 系统。"
        "当前最有说服力的点，恰恰是：在能力不完整时选择硬关闭高风险出口，而不是假装闭环已经安全。",
        bold=True,
    )

    # PPT
    add_heading(doc, "二、配套 PPT 大纲（7 页）", 1)
    ppt_pages = [
        (
            "P1 项目一句话｜约 1 min",
            [
                "面向 GitHub 改动的 AI Coding Agent（Vue3 + FastAPI + LangGraph）",
                "核心价值：编排 + HITL + 安全门禁，而非只生成代码",
                "诚实边界：真实发布当前硬关闭",
                "建议图：产品截图（任务时间线）",
                "口头：场景引入，别吹商业价值",
            ],
        ),
        (
            "P2 架构总览｜约 2 min",
            [
                "前端 SSE ↔ FastAPI ↔ LangGraph ↔ LLM / Workspace / GitHub / Redis",
                "Checkpointer ≠ RedisService",
                "建议图：架构 Mermaid",
                "口头：走一遍请求链路",
            ],
        ),
        (
            "P3 主流程状态机｜约 2 min",
            [
                "分析→计划审批→编码→Diff 审批→测试→发布审批→Publish",
                "LLM / 确定性 / 人工 三分",
                "建议图：流程图；可截 graph.py 节点列表",
                "口头：强调 Diff 前冻结",
            ],
        ),
        (
            "P4 HITL 与 Checkpoint｜约 2 min",
            [
                "interrupt / thread_id / Command(resume)",
                "Redis Stack vs MemorySaver；无耐久则 503",
                "建议图：审批时序；可贴 await_*_node",
                "口头：区分代码支持与跨重启已验证",
            ],
        ),
        (
            "P5 安全：哈希 + 写门禁 + 测试门禁｜约 2.5 min",
            [
                "approved_snapshot + content_hash",
                "permit_github_write",
                "evaluate_test_gate 恒拒 passed",
                "建议图：发布门禁漏斗",
                "口头：为什么不信模型",
            ],
        ),
        (
            "P6 技术难点页｜约 2 min",
            [
                "「批准内容」与「实际上传字节」一致性",
                "解决方案与 trade-off",
                "建议图：workspace 目录 / publish 拒绝审计",
            ],
        ),
        (
            "P7 现状与下一步｜约 1.5 min",
            [
                "已完成 / 未验证 / 硬关闭",
                "三个改进方向",
                "口头：收束金句",
            ],
        ),
    ]
    for title, bullets in ppt_pages:
        add_heading(doc, title, 2)
        add_bullets(doc, bullets)

    # 追问
    add_heading(doc, "三、技术面试官追问清单", 1)
    qa = [
        (
            "1. 为什么选 LangGraph 而不是普通 Python 或 AgentExecutor？",
            "考察：长流程与 HITL 选型。",
            "普通脚本难暂停恢复；AgentExecutor 偏工具循环。我需要计划/Diff/发布多停点，以及测试失败回编码（MAX_REPAIR=2）。"
            "LangGraph 把这些做成节点和边，再用 Checkpointer 挂起。",
            "无需额外核查。",
        ),
        (
            "2. 状态、节点、边各干什么？",
            "考察：StateGraph 基本功。",
            "State 是共享黑板（phase、plan、hash、三道 approved）；节点做一步并返回增量 dict；条件边读状态决定下一节点。",
            "无需额外核查。",
        ),
        (
            "3. interrupt 与 checkpoint 关系？",
            "考察：HITL 机制深度。",
            "interrupt 是挂起点；没有 checkpointer 无法可靠恢复。审批 API 用同一 thread_id 取 checkpoint，再 Command(resume=decision) 从断点继续。",
            "无需额外核查。",
        ),
        (
            "4. thread_id 如何关联任务？",
            "考察：会话主键。",
            "新对话生成 thread_id，写入 LangGraph config，同时 task_run_service 建任务；审批、取消、状态查询都靠它对齐。",
            "无需额外核查。",
        ),
        (
            "5. 后端重启后审批能否恢复？",
            "考察：诚实度 + 持久化。",
            "仅 checkpointer_mode==redis 时 hitl_resume_guaranteed。MemorySaver 进程没了状态就没了；无耐久则默认拒新 HITL。",
            "现场看 /health persistence；verify_hitl_restart.py 是否跑过。",
        ),
        (
            "6. Redis 不可用会发生什么？",
            "考察：降级模式。",
            "Checkpointer 降级 Memory；RedisService 可内存 fallback。新 HITL 默认 503；发布因幂等不可用被拒。",
            "看 redis_fallback 实际值。",
        ),
        (
            "7. SSE 断开后如何处理？",
            "考察：流式工程。",
            "disconnect → request_cancel；非 EventSource 自动重连。恢复靠任务状态与审批 resume。",
            "前端是否有 task 状态轮询补洞。",
        ),
        (
            "8. 如何防止错误代码直接发布？",
            "考察：纵深防御。",
            "计划未确认不编码；只改工作区；Diff 绑哈希；测试门禁硬关；发布需三门 approved；写 API 必须 WritePermit。",
            "无需额外核查。",
        ),
        (
            "9. Diff 快照和哈希解决什么？",
            "考察：内容一致性 / TOCTOU。",
            "防止展示/批准内容与推送内容不一致；work 漂移只审计；快照被改则失效审批。",
            "无需额外核查。",
        ),
        (
            "10. 为什么测试门禁不能信模型的 passed？",
            "考察：安全意识。",
            "状态可伪造；无隔离执行器证明前，对 passed=True 返回 untrusted_passed_no_secure_executor。有意硬关闭。",
            "无需额外核查。",
        ),
        (
            "11. 两个审批请求并发会怎样？",
            "考察：并发与幂等。",
            "先到者 resolve；后到者 409。发布侧还有 idempotency key。勿编造压测数字。",
            "极端并发原子性未见压测数据。",
        ),
        (
            "12. 如何防重复提交审批或重复写？",
            "考察：幂等设计。",
            "resolved 后拒再 resume；发布用 Redis 幂等锁；无 Redis 直接拒发。",
            "无需额外核查。",
        ),
        (
            "13. 当前测试能证明什么、不能证明什么？",
            "考察：测试边界诚实度。",
            "能证明路由、门禁拒绝、快照完整性、写工具门闩、进程内 HITL。不能证明生产 DeepSeek、Redis Stack 永远可用、真实 PR、全浏览器 E2E。",
            "CI 是否跑全量、Redis B 类是否 skip。",
        ),
        (
            "14. DeepSeek 怎么接？失败怎么办？",
            "考察：模型接入。",
            "ChatOpenAI + base_url/api_key/model；流式失败有 retry/规则兜底；无单独 DeepSeek SDK。",
            "现场 Key 与限流表现。",
        ),
        (
            "15. 当前最大技术债？",
            "考察：自我认知。",
            "缺可信测试执行器导致发布闭合；README 叙事偏旧 ReAct；微服务/RAG 等「代码在、证据不齐」。优先补执行器与 HITL 耐久验收。",
            "无需额外核查。",
        ),
        (
            "16.（加分）旧 ReAct 的 tools_approved 为何拒写？",
            "考察：安全收敛。",
            "写权限收拢到 Publish Agent 单一出口，避免工具循环绕过三门审批与哈希协议。",
            "可读 tools_approved_node。",
        ),
    ]
    for q, exam, ans, check in qa:
        add_heading(doc, q, 2)
        add_para(doc, f"面试官想考察：{exam}")
        add_para(doc, f"参考回答（约 30–60 秒）：{ans}")
        add_para(doc, f"需核查：{check}")

    # 多时长
    add_heading(doc, "四、三种时长讲解版本", 1)
    add_heading(doc, "3 分钟电梯版", 2)
    add_para(
        doc,
        "我做了一个 GitHub Coding Agent：Vue3 + FastAPI + LangGraph。解决「模型改仓库」缺规划、缺状态、缺人工闸门的问题。"
        "写任务走分析→计划审批→工作区改码→Diff 审批（冻结 content_hash）→测试→发布审批。"
        "模型用 DeepSeek 兼容接口。不信任模型自报测试通过；无隔离执行器就硬关远程发布；GitHub 写必须 WritePermit。"
        "当前价值是把编排、HITL 和安全协议工程化，而不是已经能量产自动合 PR。",
    )
    add_heading(doc, "5 分钟精简版", 2)
    add_para(
        doc,
        "在 3 分钟基础上补三点：① SSE 推节点事件，审批用 interrupt + 同 thread_id 的 Command(resume)；"
        "② Redis 分 Checkpointer 与业务缓存，Memory 模式不能吹跨重启；"
        "③ 难点是批准内容与上传字节一致性，用 approved_snapshot + 哈希 + 测试门禁解决，并接受发布暂时闭合。"
        "下一步做可信执行器与 HITL 恢复验收。",
    )
    add_heading(doc, "15 分钟完整版", 2)
    add_para(doc, "即本文「一、完整口述稿」七部分；不要为精彩而加入「已成功自动提 PR」等无证据表述。")

    # 演示
    add_heading(doc, "五、现场演示脚本（2–3 分钟）", 1)
    add_heading(doc, "提前启动", 2)
    add_bullets(
        doc,
        [
            "配置 backend/.env（LLM_API_KEY、可选 GITHUB_TOKEN、OPERATOR_API_TOKEN）",
            "建议起 Redis Stack；本机可用 scripts/start-all.ps1（注意 8001 vs 8000）",
            "前端 npm run dev",
            "跑 scripts/demo.ps1 看 /health：重点 checkpointer / persistence / redis",
        ],
    )
    add_heading(doc, "浏览器操作（安全优先，不碰真实发布）", 2)
    add_bullets(
        doc,
        [
            "演示 A（稳）：列出 issues / 读文件——展示 SSE 节点流与工具事件",
            "演示 B（核心）：小改动需求 → 计划确认 → 可拒绝（强调仓库未改）或同意后编码",
            "若到 Diff：指出 content_hash / 冻结快照；说明当前版本不会真实建 PR",
            "若到发布审批：同意后预期被 test_gate 拒绝——这是安全演示高潮",
        ],
    )
    add_heading(doc, "环境不可用时的诚实话术", 2)
    add_table(
        doc,
        ["情况", "说法", "切换方案"],
        [
            ["无 Redis Stack", "Memory 不可跨重启；未开 ephemeral 会 503", "ALLOW_EPHEMERAL_HITL=true 仅本机声明限制"],
            ["无 LLM Key", "无法走分析/编码", "展示单测、/health、门禁代码"],
            ["无 GitHub Token", "读写 API 失败", "只演示工作区 Diff 与门禁拒绝"],
            ["想「出 PR」", "不要现场真实发布", "展示 evaluate_test_gate 与 security 测试"],
        ],
    )
    add_heading(doc, "演示禁忌", 2)
    add_bullets(
        doc,
        [
            "不要对真实生产仓库点发布并期望出 PR",
            "不要把 MemorySaver 说成 Redis 持久化",
            "不要按 README 验收路径承诺「同意就出 PR」",
        ],
    )

    add_heading(doc, "六、面试表达建议", 1)
    add_bullets(
        doc,
        [
            "主叙事用多 Agent 三道 HITL；旧 ReAct 当旁路",
            "把「发布硬关闭」讲成设计决策",
            "个人贡献如实表述，勿默认全部独立完成",
            "被问性能数字：直接说没有压测数据，讲设计和验证边界",
        ],
    )

    path = OUT / "面试项目介绍-完整版.docx"
    doc.save(path)
    return path


def build_oral():
    doc = Document()
    cover(doc, "GitHub Agent · 15 分钟口述稿", "可直接练习的口语稿（第一人称）")
    sections = [
        (
            "第一部分：项目背景与核心价值（约 1 分钟）",
            [
                "我做的这个项目叫 GitHub Agent，本质是一个面向仓库改动的 AI Coding Agent Web 应用。",
                "传统「聊天生成代码」缺任务拆解、状态管理和高风险写操作的人工闸门。",
                "目标是把需求分析→规划→人工确认→工作区改码→Diff 审查→测试门禁→发布决策串成状态机。",
                "技术栈：Vue 3、FastAPI、LangGraph、DeepSeek（OpenAI 兼容）。",
                "诚实点：真实远程发布故意硬关闭——还没有可信隔离测试执行器。",
            ],
        ),
        (
            "第二部分：整体架构（约 2 分钟）",
            [
                "四层：Vue 前端（SSE）→ FastAPI → LangGraph 编排 → LLM / Workspace / GitHub / Redis。",
                "调用链：/api/chat/stream → assert_hitl_task_allowed → stream_agent → astream_events → 前端事件。",
                "审批：/api/approval/{tid} → Command(resume=decision) 从 checkpoint 继续。",
                "提醒：Checkpointer 与 RedisService 是两套职责，面试别混。",
            ],
        ),
        (
            "第三部分：核心业务流程（约 3 分钟）",
            [
                "classify → jev → analysis → planning → 计划 HITL → coding（仅 work/）→ Diff HITL（冻结 hash）→ testing（not_run）→ 发布 HITL → publish（门禁硬关）。",
                "LLM 理解与改码；边与门禁确定性；三处人工决策；不信模型 passed。",
            ],
        ),
        (
            "第四部分：关键技术设计（约 4 分钟）",
            [
                "LangGraph：长流程 + 条件路由 + 可挂起。",
                "HITL：interrupt + thread_id + checkpoint + Command(resume)；跨重启看 Redis Stack。",
                "SSE：节点级可见；断线 cancel，非自动重连。",
                "Redis：耐久 Checkpointer vs 业务缓存；Memory 不能吹跨重启。",
                "安全：快照哈希 + WritePermit + evaluate_test_gate 恒拒。",
            ],
        ),
        (
            "第五部分：设计难点（约 2 分钟）",
            [
                "难点：人批准的内容必须等于实际上传字节。",
                "freeze_approved_snapshot + content_hash + 发布只推内存验证快照；测试无执行器则硬关。",
                "trade-off：发布闭合换安全边界清晰。",
            ],
        ),
        (
            "第六部分：现状与规划（约 1 分钟）",
            [
                "已完成状态机/HITL/哈希/门禁；未充分证明跨重启与 E2E；未做可信执行器。",
                "下一步：隔离执行器、HITL 恢复验收、任务观测。",
            ],
        ),
        (
            "第七部分：总结（约 1 分钟）",
            [
                "不是聊天生成器，而是编排 + 状态 + 人工决策 + 工程安全约束的 Coding Agent。",
                "金句：能力不完整时硬关闭高风险出口，比假装闭环更经得起追问。",
            ],
        ),
    ]
    for title, paras in sections:
        add_heading(doc, title, 1)
        for p in paras:
            add_para(doc, p)
    path = OUT / "面试口述稿-15分钟.docx"
    doc.save(path)
    return path


def build_ppt():
    doc = Document()
    cover(doc, "GitHub Agent · PPT 大纲", "建议 7 页 · 以图和关键设计为主")
    pages = [
        ("第 1 页｜项目一句话（1 min）", ["定位：AI Coding Agent Web", "价值：编排+HITL+安全", "边界：发布硬关闭", "配图：时间线截图"]),
        ("第 2 页｜架构总览（2 min）", ["四层关系", "SSE 调用链", "Checkpointer≠RedisService", "配图：架构图"]),
        ("第 3 页｜主流程状态机（2 min）", ["分析→三道审批→编码/测试/发布", "LLM/确定/人工三分", "配图：流程图"]),
        ("第 4 页｜HITL 与 Checkpoint（2 min）", ["interrupt / resume / thread_id", "Redis Stack vs Memory", "配图：审批时序"]),
        ("第 5 页｜安全门禁（2.5 min）", ["snapshot + hash", "WritePermit", "test_gate 恒拒", "配图：门禁漏斗"]),
        ("第 6 页｜技术难点（2 min）", ["批准内容=上传字节", "trade-off：发布闭合", "配图：workspace 结构"]),
        ("第 7 页｜现状与下一步（1.5 min）", ["已完成/未验证/硬关闭", "执行器、HITL 验收、观测", "收束金句"]),
    ]
    for t, bs in pages:
        add_heading(doc, t, 1)
        add_bullets(doc, bs)
    path = OUT / "面试PPT大纲.docx"
    doc.save(path)
    return path


def build_qa():
    doc = Document()
    cover(doc, "GitHub Agent · 技术追问清单", "16 题 · 考察点 + 参考回答 + 核查提示")
    items = [
        ("为什么选 LangGraph？", "长流程多停点+可恢复；优于脚本/AgentExecutor。"),
        ("状态/节点/边职责？", "黑板；一步增量；条件路由。"),
        ("interrupt 与 checkpoint？", "挂起靠 interrupt，恢复靠 checkpoint + Command(resume)。"),
        ("thread_id？", "会话主键，对齐图配置、任务记录、审批。"),
        ("重启后审批能否恢复？", "仅 Redis Stack AsyncRedisSaver 保证；Memory 不行。"),
        ("Redis 不可用？", "Memory 降级；新 HITL 默认 503；发布拒幂等。"),
        ("SSE 断开？", "cancel；非自动重连；靠状态/审批恢复。"),
        ("防错误发布？", "三门审批+哈希+门禁+WritePermit。"),
        ("Diff 快照/哈希？", "批准内容与推送字节一致。"),
        ("为何不信 passed？", "可伪造；无执行器证明则 untrusted_passed。"),
        ("并发审批？", "先 resolve 后 409；发布幂等锁。"),
        ("防重复写？", "审批 resolved + Redis idempotency。"),
        ("测试能证明什么？", "门禁/路由/快照；不能证明生产 E2E PR。"),
        ("DeepSeek 接入？", "ChatOpenAI 兼容 base_url；失败 retry/规则兜底。"),
        ("最大技术债？", "缺可信执行器；README 叙事滞后。"),
        ("tools_approved 拒写？", "写权限收敛到 Publish Agent。"),
    ]
    for i, (q, a) in enumerate(items, 1):
        add_heading(doc, f"{i}. {q}", 2)
        add_para(doc, a)
    add_para(doc, "完整展开版见《面试项目介绍-完整版.docx》第三章。", bold=True)
    path = OUT / "面试追问清单.docx"
    doc.save(path)
    return path


def build_demo_and_versions():
    doc = Document()
    cover(doc, "GitHub Agent · 演示脚本与多时长版本", "3/5/15 分钟版本 + 现场演示")

    add_heading(doc, "一、3 分钟电梯版", 1)
    add_para(
        doc,
        "GitHub Coding Agent：Vue3 + FastAPI + LangGraph。解决模型改仓库缺规划/状态/人工闸门的问题。"
        "主路径：分析→计划审批→工作区改码→Diff 审批（content_hash）→测试→发布审批。"
        "DeepSeek 兼容接入。不信模型 passed；无隔离执行器硬关远程发布；写操作需 WritePermit。"
        "价值是编排+HITL+安全协议工程化，不是已能量产自动合 PR。",
    )
    add_heading(doc, "二、5 分钟精简版", 1)
    add_para(
        doc,
        "补三点：SSE 节点事件与 interrupt/resume；Redis 双角色（Checkpointer vs 业务）；"
        "批准内容与上传字节一致性（snapshot+hash+门禁），接受发布暂时闭合。"
        "下一步：可信执行器、HITL 恢复验收。",
    )
    add_heading(doc, "三、15 分钟完整版", 1)
    add_para(doc, "见《面试口述稿-15分钟.docx》或完整版第一章。")

    add_heading(doc, "四、现场演示脚本（2–3 分钟）", 1)
    add_heading(doc, "启动检查", 2)
    add_bullets(
        doc,
        [
            "start-all.ps1 + 前端；demo.ps1 查 /health",
            "看 persistence/checkpointer，勿假装 Redis 成功",
        ],
    )
    add_heading(doc, "演示步骤", 2)
    add_bullets(
        doc,
        [
            "A：只读问答，展示 SSE 时间线",
            "B：小改动 → 计划确认（可拒绝）→ Diff（讲 hash）",
            "高潮：发布被 test_gate 拒绝 = 安全设计",
        ],
    )
    add_heading(doc, "降级话术", 2)
    add_bullets(
        doc,
        [
            "无 Redis Stack：声明 Memory 限制，或 ephemeral 本机演示",
            "无 LLM/GitHub：展示单测与门禁代码",
            "禁止现场真实生产发布",
        ],
    )
    path = OUT / "面试演示与多时长版本.docx"
    doc.save(path)
    return path


def main():
    paths = [
        build_full(),
        build_oral(),
        build_ppt(),
        build_qa(),
        build_demo_and_versions(),
    ]
    for p in paths:
        print(p)


if __name__ == "__main__":
    main()
