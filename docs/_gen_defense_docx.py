# -*- coding: utf-8 -*-
"""生成《项目答辩逐字稿.docx》，内容与同目录 Markdown 对齐。"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent / "项目答辩逐字稿.docx"


def set_run_font(run, *, size=12, bold=False, color=None, name="微软雅黑", italic=False):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = name
    if color:
        run.font.color.rgb = color
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name)


def style_doc(doc: Document):
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    # 页码
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    set_run_font(run, size=9, color=RGBColor(0x71, 0x80, 0x96))
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr)
    run._r.append(fld_char_end)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    sizes = {1: 16, 2: 14, 3: 12}
    for run in p.runs:
        set_run_font(run, size=sizes.get(level, 12), bold=True, name="微软雅黑", color=RGBColor(0x1A, 0x36, 0x5D))
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(8)
    return p


def add_para(doc, text, *, bold=False, size=12, first_line=True, tip=False, space_after=8):
    p = doc.add_paragraph()
    run = p.add_run(text)
    color = RGBColor(0x6B, 0x46, 0x0C) if tip else RGBColor(0x1A, 0x20, 0x2C)
    set_run_font(run, size=size, bold=bold, color=color, italic=tip)
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if first_line and not tip:
        pf.first_line_indent = Cm(0.74)
    return p


def add_tip(doc, text):
    return add_para(doc, "【答辩提示】" + text, tip=True, first_line=False, size=10.5)


def add_code(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=9.5, name="Consolas", color=RGBColor(0x2D, 0x37, 0x48))
    # 东亚回退
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), "微软雅黑")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(0.4)
    return p


def add_bullets(doc, items, *, size=11.5):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(it)
        set_run_font(run, size=size)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.35


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, size=10, bold=True, color=RGBColor(0x1A, 0x36, 0x5D))
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            set_run_font(run, size=9.5)
    doc.add_paragraph()


def cover(doc: Document):
    for _ in range(3):
        doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("GitHub Agent")
    set_run_font(r, size=28, bold=True, color=RGBColor(0x1A, 0x36, 0x5D))

    t2 = doc.add_paragraph()
    t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = t2.add_run("项目答辩逐字稿")
    set_run_font(r2, size=22, bold=True, color=RGBColor(0x2B, 0x6C, 0xB0))

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = s.add_run("毕业设计 / 实习答辩 / 项目汇报 / 技术面试")
    set_run_font(r3, size=12, color=RGBColor(0x4A, 0x55, 0x68))

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r4 = note.add_run("依据 github-agent 工作区真实代码整理 · 生成日期 2026-03-26")
    set_run_font(r4, size=10.5, color=RGBColor(0x71, 0x80, 0x96))

    note2 = doc.add_paragraph()
    note2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r5 = note2.add_run("正文可直接照读；「答辩提示」「待补充确认」为备课备注")
    set_run_font(r5, size=10, color=RGBColor(0x71, 0x80, 0x96))

    doc.add_page_break()


def build():
    doc = Document()
    style_doc(doc)
    cover(doc)

    # —— 概述 ——
    add_heading(doc, "一、项目概述", 1)
    add_para(
        doc,
        "GitHub Agent（企业级智能研发助手）面向软件研发场景：用户用自然语言描述查 Issue、搜代码、改文件、提 PR 等需求，"
        "系统通过多阶段 Agent 流水线完成分析、规划、隔离区编码、人工审批与受控发布。",
    )
    add_para(
        doc,
        "一句话定位：不是自动往主干乱改代码的机器人，而是「可读可写、可中断、可审批、可审计」的 GitHub 研发 Agent。"
        "系统主体位于 github-agent/：Vue3 前端、FastAPI + LangGraph 后端、可选微服务与 MCP，以及 Redis / Postgres+pgvector。",
    )

    add_heading(doc, "二、项目背景与用户痛点", 1)
    add_para(
        doc,
        "日常开发中，大量操作是围绕 GitHub 仓库的读写链条：列 Issue、读文件、按意见改几处代码、开分支、提交、创建 Pull Request。"
        "步骤本身不难，但路径长、上下文切换多，且写操作一旦做错成本高——误改主干、误改 CI 或密钥文件、重复提 PR。",
    )
    add_table(
        doc,
        ["痛点", "具体表现", "为何值得解决"],
        [
            ["操作链条长", "浏览器、IDE、CLI、网页来回切换", "简单修改也要多步，认知负担高"],
            ["模型直连写仓库风险高", "缺少闸门时可能改错分支/敏感文件", "一次幻觉可能导致事故"],
            ["对话上下文易撑爆", "多轮工具调用后 messages 膨胀", "成本上升、协议报错、效果变差"],
            ["代码问答缺仓库感知", "只靠模型「猜」结构", "回答难落到真实路径"],
            ["审批与执行脱节", "无法在系统内确认 Diff", "责任边界不清"],
        ],
    )
    add_para(
        doc,
        "目标用户包括需要频繁读写 GitHub 的开发者与实习工程师，以及需要演示 Agent + HITL 工程能力的学生与求职者。"
        "本系统定位为演示级工程实现，代码未证明已规模化上线。真实用户故事请自行补充（待补充确认）。",
        first_line=True,
    )

    add_heading(doc, "三、项目目标与整体解决方案", 1)
    add_bullets(
        doc,
        [
            "自然语言驱动 GitHub 研发闭环：读 → 搜 → 改 → 审 → 发",
            "高危写操作必须人工确认：计划 / Diff / 发布分层门禁",
            "隔离编码：Coding Agent 只改本地工作区，默认不写远程",
            "可恢复中断：LangGraph Checkpointer + interrupt/resume",
            "可观测与可降级：限流、幂等、RAG 本地回退、测试门禁硬关闭",
        ],
    )
    add_para(
        doc,
        "写任务主路径：意图分类 → Jev 编排 → 分析 → 规划 →【确认计划】→ 工作区编码 →【确认 Diff】→ Testing →【发布审批】→ Publish。"
        "读/问答路径：上下文压缩 → ReAct 选工具 → 护栏 →（必要时）审批。",
    )

    add_heading(doc, "四、技术栈与系统架构", 1)
    add_heading(doc, "4.1 技术栈职责", 2)
    add_table(
        doc,
        ["层级", "技术", "在本项目中的职责"],
        [
            ["前端", "Vue3 / Router / Pinia / Vite / TS", "页面、Token、SSE 流式展示"],
            ["前端增强", "Monaco / axios / fetch SSE", "Diff 编辑；直连后端避免代理缓冲"],
            ["后端", "FastAPI / Pydantic", "chat、approval、repos、webhook、用量"],
            ["编排", "LangGraph / LangChain", "状态机、interrupt、工具调用"],
            ["存储", "Redis", "Checkpointer、令牌桶、幂等、会话"],
            ["检索", "Postgres+pgvector / BM25 / CE", "代码分块与混合精排"],
            ["安全", "Operator Token / 护栏 / 写门闩", "认证与 GitHub 写入许可"],
        ],
    )
    add_heading(doc, "4.2 架构说明", 2)
    add_para(
        doc,
        "前端经 SSE/HTTPS 访问 FastAPI；FastAPI 驱动 LangGraph；图侧连接 Redis 检查点、隔离工作区、GitHub API 与 pgvector。"
        "MCP Server 复用同一套工具，可供 Cursor 等客户端调用。微服务模式下 Gateway 转发 Agent/Repo，业务仍主要在 backend/app。",
    )
    add_heading(doc, "4.3 选型理由", 2)
    add_para(
        doc,
        "选 LangGraph 是因为需要可暂停、可恢复的状态机，而不是一次性脚本；选 Vue3+SSE 是因为长任务要实时看见节点与审批弹窗；"
        "Redis 三重职责让中断恢复、配额和幂等共用一套基础设施；隔离工作区加写门闩把「生成 Diff」与「写远程」拆开，符合安全默认拒绝。",
    )

    add_heading(doc, "五、核心功能讲解", 1)
    add_heading(doc, "功能 1：自然语言驱动 GitHub 读写", 2)
    add_para(
        doc,
        "用户选择 owner/repo 并输入自然语言，经操作员 Token 认证后进入流式对话。系统做意图分类，读路径走压缩与 ReAct 工具，"
        "返回 SSE 事件与最终回答。证据：api/chat.py、agent/graph.py、frontend api streamSSE。",
    )
    add_heading(doc, "功能 2：多 Agent 写任务流水线", 2)
    add_para(
        doc,
        "写任务依次经过 Analysis、Planning、计划 HITL、Coding（仅工作区）、Diff HITL、Testing、发布 HITL、Publish。"
        "价值是每步可解释、可拦截，编码与发布解耦。证据：graph.py 节点装配，agents/coding_agent.py、publish_agent.py、testing_agent.py。",
    )
    add_heading(doc, "功能 3：三层人工审批 HITL", 2)
    add_para(
        doc,
        "计划确认、Diff 确认、发布确认通过 interrupt 挂起；approval_gate 校验 pending 未过期且任务状态匹配后才 resume；"
        "拒绝则 finalize 且不执行写副作用。证据：nodes.py 中 await_*、services/approval_gate.py、ApprovalModal 等前端组件。",
    )
    add_heading(doc, "功能 4：代码 RAG 混合检索", 2)
    add_para(
        doc,
        "代码友好分隔符分块，embedding 入 pgvector 或本地 JSON 降级；BM25+向量 RRF，可选 Cross-Encoder 精排；工具 search_codebase。"
        "证据：rag_service.py、transformer_rerank.py、init_db.sql 表 code_chunks。",
    )
    add_heading(doc, "功能 5：六层护栏与统一写门禁", 2)
    add_para(
        doc,
        "拦截危险分支、敏感路径、超规模变更；高危操作进审批；幂等与 SHA 乐观锁；write_gate 无 permit 拒绝 GitHub 写。"
        "证据：guardrails.py、write_gate.py。",
    )

    add_heading(doc, "六、技术难点与解决方案", 1)

    add_heading(doc, "难点 1：可中断可恢复的 HITL 状态机", 2)
    add_para(
        doc,
        "难点在于写操作中途等人审批时不能丢状态，也不能盲 resume 旧检查点。通过 compile(checkpointer)+interrupt，"
        "审批 API 走 approval_gate 默认拒绝，规划节点命名为 planner 以避免与 state 字段 plan 冲突。"
        "效果：同意从断点继续，拒绝无写副作用。体现对 Agent「状态、副作用、恢复」的工程理解。",
    )
    add_heading(doc, "难点 2：隔离工作区与内容哈希", 2)
    add_para(
        doc,
        "LLM 修改必须先给人看 Diff，发布必须是批准时的那份内容。采用 base 基线 + work 可写副本，固化 approved_snapshot 与 content_hash，"
        "Publish 校验哈希并在 permit_github_write 内写远程。体现安全证据链思维。",
    )
    add_heading(doc, "难点 3：上下文工程", 2)
    add_para(
        doc,
        "Checkpointer 保留完整轨迹便于恢复审计；进模前 compact_context 按 token 预算裁剪/摘要，并保持 tool_calls 与 ToolMessage 配对。"
        "实现见 context_engine.py。",
    )
    add_heading(doc, "难点 4：测试门禁诚实硬关闭", 2)
    add_para(
        doc,
        "无可信沙箱时 Testing 返回 unavailable/not_run；evaluate_test_gate 对伪造 passed=True 一律拒绝。"
        "不把未完成能力包装成已完成，反而体现工程成熟度。",
    )
    add_heading(doc, "难点 5：混合检索与配额保护", 2)
    add_para(
        doc,
        "BM25+向量 RRF 与可选精排解决标识符检索弱点；Redis Lua 令牌桶原子限流，贴近 GitHub 配额，避免 INCR+EXPIRE 超发。",
    )

    add_heading(doc, "七、项目亮点", 1)
    add_bullets(
        doc,
        [
            "三门 HITL + 写门闩：生成修改与写远程切开，默认拒绝写入（工程特色，非「用了框架」本身）",
            "隔离工作区真实 Diff：人审 unified Diff，不是模型一句「我改好了」",
            "上下文压缩与检查点分离：成本与可恢复性兼顾",
            "六层护栏含路径穿越细节：工具后、副作用前拦截",
            "可降级闭环：MemorySaver、本地 RAG、伪向量、微服务熔断、LLM 用量表",
        ],
    )

    add_heading(doc, "八、测试与验证", 1)
    add_para(
        doc,
        "backend/tests 覆盖意图与路由、护栏与安全阶段、审批与 HITL、上下文引擎、工作区与多 Agent、精排、MCP、网关、用量与评测等。"
        "请在答辩前执行 pytest 并记录真实结果，本文不编造通过率。手工路径：配置 .env → 启动 → Token 登录 → 只读任务 → 受控写任务（注意测试门禁与 Token 权限）。",
    )

    add_heading(doc, "九、项目不足与未来规划", 1)
    add_para(
        doc,
        "不足：可信测试执行器未落地；认证为共享 Operator Token；部分能力可选/可视；微服务增加运维复杂度。"
        "规划（勿说成已完成）：沙箱证明链、会话 Cookie 登录、细粒度 ACL、扩大评测与 CI 门禁。",
    )

    doc.add_page_break()
    add_heading(doc, "十、完整答辩逐字稿（可直接照读）", 1)
    add_tip(doc, "建议时长 8～12 分钟。可按时间删减部分段落。")

    add_heading(doc, "第一部分：开场与项目介绍", 2)
    add_para(
        doc,
        "各位老师 / 各位评委大家好。我汇报的项目名称是 GitHub Agent，中文可以叫「企业级智能研发助手」。",
    )
    add_para(
        doc,
        "这个项目要解决的问题很具体：我们能不能用自然语言，在 GitHub 上完成「读 Issue、搜代码、改文件、建分支、提交、提 Pull Request」"
        "这一整条研发链路，同时又保证高风险写操作必须经过人的确认，而不是让大模型直接改主干。",
    )
    add_para(
        doc,
        "一句话概括：它是一个带人工审批、带隔离工作区、可中断可恢复的 GitHub 研发 Agent，而不是一个只会聊天的网页。",
    )

    add_heading(doc, "第二部分：项目背景与用户痛点", 2)
    add_para(
        doc,
        "在实际开发里，很多修改其实不大，但操作链条很长。开发者往往要在 IDE、终端、GitHub 网页之间来回切换：先找 Issue，再定位文件，再改代码，再开分支，再提 PR。"
        "更麻烦的是，如果直接把大模型接到 GitHub 写接口上，模型一旦幻觉，就可能改错分支、改到 .env 或工作流文件，甚至重复创建 PR。",
    )
    add_para(
        doc,
        "所以这个项目的出发点不是「让 AI 看起来很聪明」，而是：在自动化提效的同时，把副作用关进可审批、可审计的流程里。",
    )
    add_tip(doc, "若有个人真实故事（课程仓库、实习仓库），在此插入约 20 秒。")

    add_heading(doc, "第三部分：项目目标与整体解决方案", 2)
    add_para(
        doc,
        "我们设定了几条硬目标。第一，自然语言要能驱动真实的 GitHub 读写。第二，写任务必须拆成多阶段：先分析、再规划、再编码、再测试门禁、最后发布。"
        "第三，编码只能发生在本地隔离工作区，远程写入必须拿到统一的写许可。第四，审批过程中系统可以挂起，人点同意后再从检查点恢复。",
    )
    add_para(
        doc,
        "整体上，用户从前端发起对话，后端用 LangGraph 跑状态机，中间用 Redis 保存检查点，用护栏拦截危险操作，用工作区生成 Diff，"
        "最后在人确认后由 Publish Agent 受控调用 GitHub API。",
    )

    add_heading(doc, "第四部分：系统架构与技术选型", 2)
    add_para(
        doc,
        "前端使用 Vue 3 和 TypeScript，通过 SSE 实时展示 Agent 节点推进，并提供计划确认、Diff 审阅和发布审批界面。"
        "后端是 FastAPI，核心编排是 LangGraph。Redis 同时承担检查点、限流和幂等；Postgres 加 pgvector 做代码检索；可选接入 Langfuse 做链路追踪。",
    )
    add_para(
        doc,
        "为什么选 LangGraph？因为我们需要的不是一次性脚本，而是一张带条件边、能 interrupt、能 resume 的状态机。"
        "节点只返回状态增量，消息列表用 reducer 合并，避免并行覆盖。另外有一个实现细节：规划节点不能叫 plan，因为会和状态字段冲突，所以代码里叫 planner——"
        "这也能说明我们是按框架约束认真建模的。",
    )
    add_para(
        doc,
        "一次写任务的主路径是：意图分类，到 Jev 编排决策，再进入仓库分析、规划、等待计划确认、工作区编码、等待 Diff 确认、Testing、等待发布审批、Publish，最后汇总。"
        "读任务则更多走上下文压缩和 ReAct 工具循环，并在工具后过护栏。",
    )
    add_tip(doc, "建议展示：首页选仓库 → 对话页节点时间线 → 计划确认 → Diff 审阅。")

    add_heading(doc, "第五部分：核心功能与业务流程", 2)
    add_para(
        doc,
        "我重点讲三条主线。第一条是对话驱动。用户输入自然语言后，/api/chat/stream 以 SSE 推送事件。"
        "前端不用干等最终结果，可以看见分类、规划、工具调用等阶段。敏感接口需要操作员 Token，服务端未配置 Token 时默认失败关闭，避免演示环境意外裸奔。",
    )
    add_para(
        doc,
        "第二条是多 Agent 写流水线。Analysis Agent 只读分析；Planning Agent 产出结构化计划，包括打算改哪些文件；"
        "Coding Agent 根据计划在 base/work 工作区写完整文件内容并生成 unified Diff，它被明确禁止直接调 GitHub 写接口；"
        "Testing Agent 在没有可信隔离执行器时会诚实返回 unavailable，而不是用模型口头说「测试通过」；"
        "Publish Agent 只有在计划、Diff、发布都批准，并且内容哈希匹配、测试门禁允许时，才会拿到写许可去建分支和提 PR。",
    )
    add_para(
        doc,
        "第三条是代码问答。我们可以先给仓库文件建索引，检索时用 BM25 和向量做 RRF 融合，有条件再上 Cross-Encoder 精排。"
        "这样回答能落到真实代码片段，而不是凭空编造路径。",
    )

    add_heading(doc, "第六部分：技术难点与解决方案", 2)
    add_para(
        doc,
        "最难的第一点是 HITL 中断恢复。高危步骤调用 interrupt 后，状态进入 Checkpointer。"
        "人审批时，我们不是简单「继续跑图」，而是先过审批门禁：没有 pending、已过期、已拒绝、任务状态不匹配，一律拒绝 resume。"
        "这样可以避免用旧检查点绕过当前业务状态。",
    )
    add_para(
        doc,
        "第二点是隔离工作区与发布哈希。人批准的是某一版 Diff 对应的快照。发布时校验 content_hash，"
        "并用上下文变量形式的 WritePermit 临时打开 GitHub 写能力，用完即收回。编辑器、普通工具默认没有这个许可。",
    )
    add_para(
        doc,
        "第三点是上下文工程。检查点里保留完整消息便于恢复和审计；真正进大模型前，由 compact 节点按 token 预算裁剪或摘要，"
        "并保持工具调用与工具结果配对，避免协议层报错。",
    )
    add_para(
        doc,
        "第四点是测试门禁的诚实性。我们明确知道宿主机直接执行仓库命令不安全，所以当前阶段 Testing 不强行跑命令。"
        "即便有人把状态写成 passed，Publish 侧也会因为「没有安全执行器证明」而拒绝。"
        "本地调试可以显式打开不信任旁路，但这是配置开关，不是默认行为。",
    )
    add_para(
        doc,
        "第五点是护栏与配额。六层护栏覆盖危险分支、敏感路径、变更规模、强制审批、幂等和冲突重试；"
        "对 GitHub 调用用 Redis Lua 令牌桶，避免高并发下超发配额。",
    )

    add_heading(doc, "第七部分：项目亮点与实际价值", 2)
    add_para(
        doc,
        "如果只用一句话讲亮点，我会说：我们把「AI 生成修改」和「AI 拥有写权限」切开了。"
        "用户价值是操作路径更短，但仍能看清计划与 Diff 再决定是否发布。技术价值是状态机、门禁、工作区、检索、限流形成闭环。"
        "工程价值是大量边界有单测和失败关闭策略，并且对未完成能力不强行伪装。",
    )

    add_heading(doc, "第八部分：测试与验证", 2)
    add_para(
        doc,
        "后端 tests 目录覆盖了意图路由、护栏、审批门禁、上下文引擎、工作区、发布错误格式、安全阶段策略、网关路由、MCP 等。"
        "答辩前我会用 pytest 跑通关键用例，并用一条端到端只读任务和一条受控写任务做人工演示。"
        "关于性能提升百分比或线上用户数，我目前没有可公开的统计，所以不在这里陈述。",
    )

    add_heading(doc, "第九部分：不足与展望", 2)
    add_para(
        doc,
        "当前最大的不足是缺少可信测试执行沙箱，因此严格模式下发布会受测试门禁限制。"
        "认证仍是共享操作员 Token，适合本地演示，不适合公网多租户。"
        "后续如果继续做，我会优先补沙箱证明链和会话型登录，再扩展评测回归。",
    )

    add_heading(doc, "第十部分：总结与致谢", 2)
    add_para(
        doc,
        "总结一下：GitHub Agent 面向「用自然语言完成 GitHub 研发操作」这一场景，用 LangGraph 状态机、多层人工审批、隔离工作区和统一写门禁，"
        "把自动化控制在可审查的边界内。我的工作重点在……（待补充确认：个人分工）。",
    )
    add_para(
        doc,
        "以上是我的汇报，感谢各位老师聆听，请各位批评指正。",
    )

    doc.add_page_break()
    add_heading(doc, "十一、核心代码定位与答辩演示建议", 1)
    add_heading(doc, "11.1 建议展示代码顺序", 2)
    add_table(
        doc,
        ["顺序", "位置", "要讲清的重点"],
        [
            ["1", "agent/graph.py → build_graph", "写流水线节点边；planner 命名"],
            ["2", "agent/state.py → AgentState", "reducer 与审批相关字段"],
            ["3", "agent/nodes.py → await_*", "interrupt 与 approval_kind"],
            ["4", "services/workspace_service.py", "base/work 隔离与 Diff"],
            ["5", "agents/coding_agent.py", "只写工作区"],
            ["6", "services/write_gate.py", "无 permit 拒绝写远程"],
            ["7", "agents/publish_agent.py", "evaluate_test_gate 拒伪造 passed"],
            ["8", "services/guardrails.py", "六层护栏"],
            ["9", "services/context_engine.py", "压缩与配对"],
            ["10", "services/approval_gate.py", "resume 默认拒绝"],
            ["11", "frontend api streamSSE", "直连与 Token"],
            ["12", "scripts/init_db.sql", "code_chunks / llm_call_records"],
        ],
    )
    add_heading(doc, "11.2 建议演示页面", 2)
    add_bullets(
        doc,
        [
            "LoginView：操作员 Token（说明不是完整 IAM）",
            "ChatView：只读指令 + SSE 时间线",
            "PlanConfirmView / ApprovalModal：计划确认",
            "ReviewView + DiffViewer：确认 Diff",
            "可选 UsageView / EvalView；条件允许则看 GitHub PR",
        ],
    )

    add_heading(doc, "十二、答辩常见问题与参考回答", 1)

    qa = [
        (
            "Q1：为什么要做这个项目？",
            "因为 GitHub 研发操作链条长，且直接让模型写仓库风险高。目标是用自然语言缩短路径，同时用 HITL 和工作区管住写风险。真实业务故事待补充确认。",
        ),
        (
            "Q2：解决了用户什么实际问题？",
            "用户原先要手动找文件、改代码、开分支、提 PR；现在在对话中发起，系统给出计划与 Diff，人确认后再发布。减少切换成本与「看不见变更就写入」的风险。",
        ),
        (
            "Q3：为什么选这些技术栈？",
            "Vue3+SSE 适合长任务可视化；FastAPI 适合异步 API；LangGraph 提供检查点与 interrupt；Redis 覆盖恢复/限流/幂等；pgvector 支撑代码检索。选型服务于可中断写流程。",
        ),
        (
            "Q4：系统整体架构如何设计？",
            "前端调 FastAPI；核心是 LangGraph 状态机；写路径走多 Agent + 工作区；读路径走工具与 RAG；Redis/Postgres 分别承担状态与检索。可选 Gateway 微服务拆分。",
        ),
        (
            "Q5：最困难的问题是什么？",
            "可恢复人工审批叠加发布内容防调包：既要 interrupt/resume，又要 approval_gate 防盲恢复，还要用 content_hash + write_gate 保证发布字节与批准一致。",
        ),
        (
            "Q6：最有价值的亮点是什么？",
            "编码与发布解耦，以及三门审批。普通 Demo 容易一键乱写；我们默认无写许可。",
        ),
        (
            "Q7：写任务具体如何实现？",
            "分类与 Jev 决定走多 Agent；分析规划后 interrupt 等计划确认；Coding 改 work 出 Diff；Diff 审批后 Testing；发布审批后 Publish 校验门禁与哈希，permit 内调 GitHub。",
        ),
        (
            "Q8：数据库为什么这样设计？",
            "code_chunks 按 repo/path/chunk 唯一并建向量索引，服务 RAG；llm_call_records 记录调用用量。会话与审批多在 Redis 与本地 .data，并非全部进 Postgres。",
        ),
        (
            "Q9：如何处理异常、权限与安全？",
            "Operator Token 失败关闭；Webhook 无 secret 默认拒绝；路径护栏；写门闩；幂等；审批过期拒绝 resume。共享 Token 局限会主动说明。",
        ),
        (
            "Q10：如何验证功能正确？",
            "以 backend/tests 为主，再加手工 SSE 验收。具体 pytest 数量待答辩前跑通后补充，不编造通过率。",
        ),
        (
            "Q11：还有哪些不足？",
            "缺可信测试沙箱；认证非多用户；公网不适合 localStorage Token；部分能力依赖可选组件。",
        ),
        (
            "Q12：下一步做什么？",
            "优先隔离执行器与测试证明，打通严格模式安全发布；其次会话登录与更细 ACL。",
        ),
        (
            "Q13：与常规实现相比值得关注什么？",
            "重点不是会不会调 GitHub API，而是状态机门禁、工作区证据链、对未完成测试能力的硬关闭。",
        ),
        (
            "Q14：用户量增大如何优化？",
            "令牌桶、微服务拆分与熔断、RAG 增量索引、检查点存储扩展。当前无大规模压测数据，待补充确认。",
        ),
        (
            "Q15：哪些工作是本人完成的？",
            "待补充确认。请按真实分工填写，切勿把未参与模块说成独立完成。",
        ),
    ]
    for q, a in qa:
        add_para(doc, q, bold=True, first_line=False, space_after=4)
        add_para(doc, "参考回答：" + a, first_line=True, space_after=10)

    add_heading(doc, "十三、待确认信息清单", 1)
    add_bullets(
        doc,
        [
            "个人姓名、学校/公司、答辩身份",
            "项目起止时间与个人分工",
            "演示仓库与 Token 脱敏方案",
            "pytest 最新运行结果",
            "是否有真实用户试用或仅自测",
            "微服务 / MCP 实操经历是否写入简历表述",
            "与 hello-agents 教程的关系说明（本稿以 github-agent 为准）",
        ],
    )
    add_para(
        doc,
        "文档结束。请以代码与可演示行为为准，勿夸大未实现能力。",
        first_line=False,
        bold=True,
    )

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
